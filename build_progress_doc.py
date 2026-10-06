from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"output\电商基础数据管理平台_需求构建进度文档.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRAY = "D9E1E8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, sz="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, bold=False, color=BLACK):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def style_paragraph(p, before=0, after=6, line=1.25, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_text(doc, text, bold_prefix=None, after=6):
    p = doc.add_paragraph()
    style_paragraph(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45 + level * 0.35)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    style_paragraph(p, after=3, line=1.15)
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(p, before=10 if level == 1 else 6, after=5, line=1.1)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, h in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            set_cell_shading(cell, "FFFFFF" if ri % 2 == 0 else PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.1,
                            align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if i == 2 and str(value).startswith("已完成"):
                set_cell_shading(cell, GREEN)
            elif i == 2 and str(value).startswith("部分"):
                set_cell_shading(cell, YELLOW)
            elif i == 2 and str(value).startswith("未完成"):
                set_cell_shading(cell, RED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout_line(doc, label, text, fill=BLUE):
    p = doc.add_paragraph()
    style_paragraph(p, after=8, line=1.25)
    r = p.add_run(label)
    set_run_font(r, size=10.5, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    return p


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5)):
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLACK


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run("电商基础数据管理平台 | 需求构建进度文档")
        set_run_font(r, size=8.5, color=RGBColor(100, 100, 100))


def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=28, after=10, line=1.0)
    r = p.add_run("电商基础数据管理平台")
    set_run_font(r, size=21, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p2, after=20, line=1.0)
    r = p2.add_run("需求构建进度与开发计划")
    set_run_font(r, size=15, bold=True, color=RGBColor(47, 84, 150))

    meta = add_table(doc, ["文档属性", "内容"], [
        ("交付对象", "全栈工程师及开发人员"),
        ("适用范围", "第一阶段基础数据管理平台、第二阶段用户行为与销售数据分析"),
        ("当前基线", "现有 Spring Boot + MyBatis + 原生 HTML/JavaScript 管理系统"),
        ("编制日期", "2026 年 9 月 17 日"),
    ], widths=[3.0, 13.2], font_size=9.5)

    add_callout_line(doc, "总体结论：", "现有系统已具备后台基础管理和商品发布的初步闭环，但尚未达到两份需求文档的完整验收标准。第一阶段需优先补齐基础规则、属性维护、商品生命周期及订单库存模块；第二阶段需要单独建设离线数仓、指标服务和分析看板。", fill=BLUE)

    add_heading(doc, "一 项目现状与进度基线", 1)
    add_text(doc, "当前代码已经形成用户、部门、角色、品牌、类别、属性组、商品发布和商品查询等基础模块。商品发布采用基本信息、规格参数、销售属性、SKU 设置、确认发布五步流程，前四步使用 Redis 暂存，最后一步写入 MySQL。", after=5)
    add_text(doc, "现有数据概览页面包含四张 ECharts 图表，其中各部门人数来自真实接口，其余季度销售额、月度销售额和产品类型占比仍为静态演示数据。因此，现有 Dashboard 只能作为页面原型，不能作为第二阶段分析成果。", after=5)
    add_text(doc, "本文件中的“已完成”表示源码中已有对应功能实现；“部分完成”表示已有主体但不满足完整验收；“未完成”表示当前没有可交付的页面、接口或数据处理实现。所有模块仍需在真实数据库、Redis 和测试数据环境中完成联调验收。", after=8)

    add_heading(doc, "二 第一阶段需求完成情况", 1)
    phase1_rows = [
        ("1", "登录与注册", "已完成", "登录、注册、退出、会话检查、业务接口登录拦截已具备；需补充环境联调和安全测试。"),
        ("2", "用户管理", "部分完成", "列表、分页、搜索、新增、编辑、删除、角色绑定已实现；缺部门树点击联动、启用/禁用、批量操作和导入。"),
        ("3", "部门管理", "部分完成", "树形展示、增删改、关联用户/子部门删除校验已实现；缺最多 4 级、按钮按节点类型控制、名称点击编辑、部门搜索，且当前允许修改父部门。"),
        ("4", "品牌管理", "部分完成", "列表、搜索、分页、Logo 上传预览、状态切换、类别关联已实现；删除未阻止已有类别或商品关联的数据。"),
        ("5", "类别管理与品牌关联", "部分完成", "多级树、一级类别新增、品牌关联已实现；缺最多 3 级、节点操作按钮规则、名称点击编辑、类别搜索、关联弹窗分页搜索及关联数据删除保护。"),
        ("6", "属性分组", "部分完成", "属性组 CRUD 和类别筛选已实现；需求要求的类别树形菜单尚未实现。"),
        ("7", "规格参数维护", "未完成", "后端存在 GoodsAttr CRUD 接口，但没有独立的前端维护页面和完整操作入口。"),
        ("8", "销售属性维护", "未完成", "发布商品时可按类别读取销售属性，但没有独立的销售属性维护页面。"),
        ("9", "商品发布与 SPU/SKU", "部分完成", "发布主流程和商品查看/删除已具备；缺 SKU 图片、优惠/售后信息、定时发布、草稿管理、SPU/SKU 编辑与上下架。"),
        ("10", "订单管理", "未完成", "没有订单实体、订单接口、状态流转、详情、备注、导出和操作日志。"),
        ("11", "库存管理", "未完成", "没有库存、库存流水、预警、盘点和订单联动。"),
        ("12", "活动管理", "未完成", "没有满减、折扣、特惠活动及商品绑定、结算联动。"),
        ("13", "会员管理", "未完成", "没有会员等级、积分、消费记录、状态管理和导出。"),
        ("14", "广告管理", "未完成", "没有广告资源、投放时间、位置、排序和前台展示联动。"),
    ]
    add_table(doc, ["编号", "模块", "状态", "当前实现与缺口"], phase1_rows,
              widths=[1.0, 3.1, 2.0, 10.1], font_size=8.8)

    add_heading(doc, "三 第一阶段遗留模块开发计划", 1)
    add_heading(doc, "3.1 基础管理模块整改", 2)
    add_bullet(doc, "用户管理：增加部门树/用户列表联动；补充用户状态字段与启用/禁用接口；保留现有分页搜索、角色绑定能力。")
    add_bullet(doc, "部门管理：后端强制限制最大 4 级；非叶子节点禁止删除；四级节点禁止新增；部门名称点击进入编辑；增加部门名称模糊搜索；根据需求决定是否禁止修改父部门。")
    add_bullet(doc, "类别管理：后端强制限制最大 3 级；非叶子节点禁止删除；三级节点禁止新增；类别名称点击编辑；增加名称搜索；删除前检查品牌、属性、商品等所有关联。")
    add_bullet(doc, "品牌管理：删除前检查类别关联、SPU/商品关联和其他业务关联，存在关联时返回明确业务提示，不得先清理关联再删除。")
    add_bullet(doc, "角色管理：补充权限节点、角色权限配置、用户权限校验和前端菜单/按钮权限控制；角色删除需检查用户关联。")

    add_heading(doc, "3.2 规格参数与销售属性维护", 2)
    add_bullet(doc, "新增规格参数管理页：按类别树筛选，支持属性名称、属性组、属性类型、值类型、候选值、启用状态和检索状态维护。")
    add_bullet(doc, "新增销售属性管理页：复用属性实体和 CRUD 能力，但固定或明确 attrType=2；支持颜色、尺码、版本等多值候选项维护。")
    add_bullet(doc, "发布商品第 2、3 步仅负责选择和填写商品属性值，不替代属性定义维护；前端页面需要提供从管理菜单进入的完整入口。")

    add_heading(doc, "3.3 商品发布与 SPU/SKU 生命周期", 2)
    add_bullet(doc, "补充 SKU 独立图片上传与预览；每个 SKU 支持图片、价格、库存、编码、状态和优惠配置。")
    add_bullet(doc, "明确草稿模型：支持草稿保存、草稿列表、继续编辑、取消草稿、过期清理，不能只依赖浏览器 localStorage。")
    add_bullet(doc, "补充 SPU 编辑、上下架、批量操作和 SKU 编辑；商品详情页提供可回填的编辑入口。")
    add_bullet(doc, "将商品发布、SKU、库存、活动、订单之间的状态和事务边界统一设计，避免发布成功但子表数据不完整。")

    add_heading(doc, "3.4 订单、库存、活动、会员、广告模块", 2)
    add_text(doc, "建议按照“订单先行、库存联动、活动结算、会员沉淀、广告投放”的顺序开发。订单和库存是交易闭环的基础，应在活动、会员和广告之前完成核心数据模型及状态日志。", after=5)
    module_rows = [
        ("订单", "订单主表、明细表、状态日志、支付/发货/完成/取消/退款状态、备注、异常处理、导出", "与用户、SPU、SKU、库存建立稳定关联"),
        ("库存", "SKU 库存初始化、增减库存、批量调库、库存流水、预警、盘点", "订单下单/支付/退款/退货触发库存变动，要求幂等"),
        ("活动", "满减、折扣、特惠、时间范围、状态、商品绑定、规则校验", "订单结算计算优惠，处理活动叠加和优先级"),
        ("会员", "会员信息、等级、积分、积分流水、消费记录、启停、筛选导出", "关联订单、活动和营销分析"),
        ("广告", "图片、链接、位置、权重、投放时间、状态、上下架", "前台只展示有效时间内且已上架资源"),
    ]
    add_table(doc, ["模块", "首版建设内容", "关键联动"], module_rows,
              widths=[2.0, 8.0, 6.2], font_size=9.0)

    add_heading(doc, "四 第二阶段数据分析建设计划", 1)
    add_callout_line(doc, "建设判断：", "第二阶段不是在现有 Dashboard 上增加几张图，而是新增一条“原始数据采集/落盘 → Hive 清洗 → 指标统计 → MySQL 落地 → SpringBoot 接口 → Vue/ECharts 看板”的数据链路。", fill=YELLOW)

    add_heading(doc, "4.1 数据处理分层", 2)
    data_rows = [
        ("原始层 ODS", "HDFS", "用户行为日志、订单/退换货、评价、客服会话原始数据", "保留原始字段和批次日期"),
        ("清洗层 DWD", "Hive", "去重、脏数据过滤、时间/设备/行为/订单状态标准化、维度补全", "按日期分区，记录清洗批次和异常数量"),
        ("汇总层 DWS", "Hive", "入口、行为趋势、转化漏斗、销售、退换货、评价、客服指标", "统一指标口径，支持日/周/月汇总"),
        ("服务层 ADS", "MySQL", "四大指标表、趋势表、榜单表、看板查询结果", "SpringBoot 只查询结构化结果，不直接扫描原始日志"),
        ("展示层", "Vue + ECharts", "行为、销售、评价、客服四大分析看板", "支持日期、维度筛选和下钻"),
    ]
    add_table(doc, ["层级", "技术位置", "职责", "交付要求"], data_rows,
              widths=[2.3, 2.5, 7.0, 5.0], font_size=8.9)

    add_heading(doc, "4.2 四大分析模块首版范围", 2)
    analysis_rows = [
        ("用户行为", "入口流量、浏览/搜索/收藏/加购/下单/支付趋势、转化漏斗、热门资源、时段热力图", "入口 UV/PV、跳出率、转化率、行为趋势、漏斗转化率、流失人数"),
        ("订单销售", "销售总览、销售趋势、品类结构、商品 TOP、订单流失、退换货、客单价分布", "订单数、成交金额、客单价、支付转化率、退换货率、退款金额"),
        ("评价舆情", "评价总览、等级分布、正负关键词、品类口碑、口碑趋势、低分商品榜", "总评价、好中差评、好评率、关键词 TOP、平均评分、追评率"),
        ("客服咨询", "会话总览、咨询类型、时段热力、咨询转化、热门咨询商品", "有效会话、响应时长、解决率、咨询后加购率和成交率"),
    ]
    add_table(doc, ["模块", "首版页面", "首版核心指标"], analysis_rows,
              widths=[2.3, 8.0, 6.5], font_size=8.8)

    add_heading(doc, "4.3 建议新增数据表", 2)
    add_text(doc, "以下为建议命名，工程师可根据现有数据库规范调整，但必须保留日期分区、指标口径、批次号和数据更新时间等字段。", after=5)
    table_rows = [
        ("业务过程表", "tbl_order、tbl_order_item、tbl_order_status_log、tbl_refund_return", "承载订单及退换货事实数据"),
        ("库存与营销", "tbl_inventory、tbl_inventory_flow、tbl_activity、tbl_activity_goods", "支撑库存和活动联动"),
        ("会员与投放", "tbl_member、tbl_member_level、tbl_points_flow、tbl_advertisement", "支撑会员和广告管理"),
        ("评价与会话", "tbl_review、tbl_customer_service_session、tbl_customer_service_message", "支撑评价和客服分析原始业务数据"),
        ("行为指标", "ads_user_behavior_daily、ads_user_funnel_daily、ads_traffic_source_daily", "行为趋势、漏斗和入口指标"),
        ("销售指标", "ads_order_sales_daily、ads_refund_daily、ads_product_sales_rank", "销售、退换货和商品榜单"),
        ("评价指标", "ads_review_sentiment_daily、ads_review_keyword_rank", "评价等级、情感和关键词"),
        ("客服指标", "ads_service_session_daily、ads_service_type_rank", "会话、类型、响应和转化"),
    ]
    add_table(doc, ["分类", "建议表名", "用途"], table_rows,
              widths=[3.0, 8.2, 5.6], font_size=8.7)

    add_heading(doc, "4.4 建议接口分组", 2)
    api_rows = [
        ("基础整改", "/User、/Dept、/Role、/Brand、/Category、/GoodsAttr", "补规则校验、权限、搜索、关联保护和维护页面所需接口"),
        ("交易模块", "/Order、/Inventory、/Activity、/Member、/Advertisement", "列表、详情、状态、操作、流水和导出接口"),
        ("数据分析", "/Analysis/behavior、/Analysis/sales、/Analysis/review、/Analysis/service", "总览、趋势、分布、榜单、漏斗、热力和时间筛选"),
        ("任务管理", "/Analysis/task、/Analysis/data-quality", "清洗任务状态、批次、数据质量和失败重跑"),
    ]
    add_table(doc, ["接口分组", "建议路径", "主要职责"], api_rows,
              widths=[3.0, 6.0, 7.8], font_size=8.8)

    add_heading(doc, "五 开发顺序与里程碑", 1)
    milestone_rows = [
        ("阶段 0", "环境与数据基线", "确认 MySQL/Redis、数据库表结构、测试账号、现有接口契约；建立 dev/test 配置", "可启动、可登录、可回滚"),
        ("阶段 1", "基础模块整改", "用户、部门、角色权限、品牌、类别、属性组/属性维护及删除保护", "第一阶段个人必做项达到验收"),
        ("阶段 2", "商品生命周期", "发布增强、草稿、SPU/SKU 编辑、上下架、SKU 图片、库存接口预留", "商品从创建到下架可追踪"),
        ("阶段 3", "订单与库存", "订单状态、订单明细、库存扣减/恢复、流水、预警、幂等", "完成交易和库存闭环"),
        ("阶段 4", "活动会员广告", "营销活动、会员积分等级、广告投放与前台联动", "完成第一阶段业务扩展模块"),
        ("阶段 5", "数据仓库与指标", "HDFS/Hive 分层、清洗任务、四类指标落 MySQL、质量校验", "每日批处理可稳定运行"),
        ("阶段 6", "分析看板与验收", "Vue/ECharts 四大看板、筛选下钻、导出、权限、性能和部署", "第二阶段需求整体验收"),
    ]
    add_table(doc, ["阶段", "主题", "主要交付", "完成标志"], milestone_rows,
              widths=[2.0, 3.2, 8.0, 4.0], font_size=8.6)

    add_heading(doc, "六 全栈工程实施要求", 1)
    add_bullet(doc, "接口规范：统一响应结构、错误码、分页结构和时间格式；新增接口必须提供请求参数、响应示例和异常场景说明。")
    add_bullet(doc, "数据一致性：订单、库存、活动、会员积分等跨表操作必须定义事务边界；库存扣减、退款恢复和批处理重跑必须幂等。")
    add_bullet(doc, "权限安全：后端接口必须做角色/权限校验，不能只依赖前端隐藏按钮；文件上传需限制类型、大小和存储路径。")
    add_bullet(doc, "数仓质量：每个指标必须登记口径、来源字段、过滤条件、更新时间和异常数据量；清洗失败支持重跑，不覆盖原始数据。")
    add_bullet(doc, "前端交互：列表、搜索、分页、表单校验、空态、错误态、加载态和操作成功后的刷新必须统一处理。")
    add_bullet(doc, "测试要求：每个模块至少包含接口测试、核心业务流程测试、边界数据测试和权限测试；商品发布、订单库存、指标批处理要有回归用例。")

    add_heading(doc, "七 验收清单", 1)
    acceptance_rows = [
        ("基础权限", "未登录不能访问业务接口；不同角色只能操作授权菜单和接口"),
        ("组织树", "部门最多 4 级、类别最多 3 级；节点新增/删除按钮符合规则；关联数据删除受保护"),
        ("属性维护", "规格参数和销售属性有独立管理入口，可按类别和属性组维护并被发布流程正确读取"),
        ("商品发布", "草稿可恢复；销售属性组合无重复无遗漏；SKU 图片、价格、库存和状态可维护；最终数据完整落库"),
        ("交易库存", "订单状态可流转；库存扣减、恢复、流水和预警准确；重复请求不会重复扣减"),
        ("业务扩展", "活动、会员、广告的配置和状态变化能够正确影响前台和订单结算"),
        ("数据分析", "四类指标按日生成并落 MySQL；看板支持日/周/月和维度筛选；图表与明细口径一致"),
        ("运维质量", "批处理有日志、失败告警和重跑；接口有权限、参数校验、分页和数据更新时间"),
    ]
    add_table(doc, ["验收项", "验收标准"], acceptance_rows,
              widths=[3.0, 14.2], font_size=9.0)

    add_heading(doc, "八 当前最先交给工程师的任务", 1)
    add_text(doc, "建议先按以下顺序建立开发任务，不要直接从第二阶段图表开始：", after=5)
    for item in [
        "建立数据库变更清单和接口契约，确认现有表结构与数据字典。",
        "完成部门、类别、品牌关联删除保护及层级规则，并补对应接口测试。",
        "增加规格参数、销售属性独立维护页面，打通属性组、类别和发布流程。",
        "补商品草稿、SKU 图片、SPU/SKU 编辑和上下架能力。",
        "设计并实现订单与库存核心表，先跑通下单、扣库存、支付、退款恢复和流水。",
        "在交易数据稳定后，再搭建 Hive 清洗、指标落 MySQL 和四大分析看板。",
    ]:
        add_bullet(doc, item)

    add_footer(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
